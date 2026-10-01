"""Geração dos arquivos .xlsx (produtos e dados irregulares)."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.product import Irregularity, ProductDetail

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

EXPORT_FILENAME = "produtos.xlsx"
HEADERS = ["id", "ean", "titulo", "preço médio"]
COLUMN_WIDTHS = [8, 16, 40, 14]

IRREGULARITIES_FILENAME = "produtos_irregulares.xlsx"
IRREGULARITIES_HEADERS = ["id", "ean", "titulo", "loja", "campo", "valor recebido", "problema"]
IRREGULARITIES_COLUMN_WIDTHS = [8, 16, 30, 22, 10, 22, 60]

TEXT_FORMAT = "@"


def _new_sheet(title: str, headers: list[str], widths: list[int]) -> tuple[Workbook, Worksheet]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for cell, width in zip(sheet[1], widths):
        sheet.column_dimensions[cell.column_letter].width = width
    return workbook, sheet


def _to_bytes(workbook: Workbook) -> bytes:
    # Gerado em memória: nada é gravado em disco.
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_products_workbook(products: list[ProductDetail]) -> bytes:
    """Gera a planilha com todos os produtos e o preço médio de cada um."""
    workbook, sheet = _new_sheet("Produtos", HEADERS, COLUMN_WIDTHS)

    for product in products:
        # O preço médio é o mesmo calculado na listagem (ProductDetail.average_price),
        # garantindo uma única fonte de verdade.
        sheet.append([product.id, product.ean, product.title, product.average_price])
        row = sheet.max_row
        # EAN gravado como texto: como número, o Excel exibiria notação científica
        # (7,89E+12) e descartaria zeros à esquerda.
        sheet.cell(row=row, column=2).number_format = TEXT_FORMAT
        # Preço numérico com 2 casas; produto sem oferta válida fica com a célula vazia.
        sheet.cell(row=row, column=4).number_format = "0.00"

    return _to_bytes(workbook)


def build_irregularities_workbook(rows: list[Irregularity]) -> bytes:
    """Gera a planilha com um dado irregular por linha."""
    workbook, sheet = _new_sheet(
        "Dados irregulares", IRREGULARITIES_HEADERS, IRREGULARITIES_COLUMN_WIDTHS
    )

    for item in rows:
        sheet.append([
            item.product_id, item.ean, item.title, item.store,
            item.field, item.received_value, item.problem,
        ])
        row = sheet.max_row
        # EAN e valor recebido como texto: o valor deve aparecer exatamente como veio da API.
        sheet.cell(row=row, column=2).number_format = TEXT_FORMAT
        sheet.cell(row=row, column=6).number_format = TEXT_FORMAT

    return _to_bytes(workbook)
